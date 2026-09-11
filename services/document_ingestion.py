from __future__ import annotations

import csv
import hashlib
import io
import re
from pathlib import Path

from services.models import DocumentChunk, DocumentRecord

CHUNK_CHARS = 1800
CHUNK_OVERLAP = 220


class DocumentParseError(ValueError):
    pass


def _clean(text: str) -> str:
    text = text.replace("\u00ad", "").replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _chunk_page(name: str, page: int, text: str, doc_hash: str) -> list[DocumentChunk]:
    text = _clean(text)
    if not text:
        return []
    chunks: list[DocumentChunk] = []
    start = 0
    index = 0
    while start < len(text):
        end = min(len(text), start + CHUNK_CHARS)
        if end < len(text):
            split = max(text.rfind("\n", start, end), text.rfind(". ", start, end))
            if split > start + 700:
                end = split + 1
        piece = text[start:end].strip()
        if piece:
            chunk_hash = hashlib.sha256(f"{doc_hash}:{page}:{index}:{piece}".encode()).hexdigest()
            chunks.append(DocumentChunk(
                id=f"{doc_hash[:10]}-p{page}-c{index}",
                document_name=name,
                page=page,
                text=piece,
                sha256=chunk_hash,
            ))
            index += 1
        if end >= len(text):
            break
        start = max(start + 1, end - CHUNK_OVERLAP)
    return chunks


def parse_document(name: str, data: bytes, mime_type: str = "") -> DocumentRecord:
    suffix = Path(name).suffix.lower()
    doc_hash = hashlib.sha256(data).hexdigest()
    pages: list[str] = []

    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            pages = [_clean(page.extract_text() or "") for page in reader.pages]
        elif suffix == ".docx":
            from docx import Document
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            pages = [_clean("\n".join(parts))]
        elif suffix in {".txt", ".md"}:
            pages = [_clean(data.decode("utf-8", "replace"))]
        elif suffix == ".csv":
            text = data.decode("utf-8-sig", "replace")
            reader = csv.reader(io.StringIO(text))
            pages = [_clean("\n".join(" | ".join(row) for row in reader))]
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
            lines: list[str] = []
            for ws in wb.worksheets:
                lines.append(f"SHEET: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    vals = ["" if v is None else str(v) for v in row]
                    if any(v.strip() for v in vals):
                        lines.append(" | ".join(vals))
            pages = [_clean("\n".join(lines))]
        else:
            raise DocumentParseError(f"Unsupported file type: {suffix or 'unknown'}")
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"Could not parse {name}: {exc}") from exc

    nonempty = [p for p in pages if p.strip()]
    if not nonempty:
        raise DocumentParseError(f"No readable text could be extracted from {name}")

    chunks: list[DocumentChunk] = []
    for page_num, page_text in enumerate(pages, start=1):
        chunks.extend(_chunk_page(name, page_num, page_text, doc_hash))

    return DocumentRecord(
        name=name,
        mime_type=mime_type or "application/octet-stream",
        sha256=doc_hash,
        page_count=max(1, len(pages)),
        text="\n\n".join(nonempty),
        chunks=chunks,
    )
